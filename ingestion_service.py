import csv
import datetime
import io
import json
import re
from html.parser import HTMLParser
from urllib.error import HTTPError, URLError
from urllib.parse import urljoin, urlparse
from urllib.request import Request, urlopen

from source_metadata_service import (
    detect_eurostat_custom_view,
    extract_eurostat_dataset_code,
    normalize_eurostat_dataset_code
)


MAX_DOWNLOAD_BYTES = 25 * 1024 * 1024
MAX_STORED_ROWS = 5000
SUPPORTED_EXTENSIONS = {
    ".csv",
    ".json",
    ".xlsx",
    ".xls"
}
EUROSTAT_JSONSTAT_ENDPOINT = (
    "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/"
    "data/{code}?format=JSON&lang=en"
)
RESOURCE_KEYWORDS = {
    "api",
    "csv",
    "data",
    "dataset",
    "download",
    "excel",
    "json",
    "table",
    "xls",
    "xlsx"
}


def normalise_value(value):

    if value == "":

        return None


    if isinstance(
        value,
        (
            datetime.datetime,
            datetime.date
        )
    ):

        return value.isoformat()


    return value


def normalise_row(row):

    return {
        str(key):
            normalise_value(
                value
            )

        for key, value
        in row.items()
    }


def detect_format(
    url,
    content_type=""
):

    path = urlparse(
        url
    ).path.lower()


    if path.endswith(
        ".csv"
    ) or "text/csv" in content_type:

        return "csv"


    if path.endswith(
        ".json"
    ) or "json" in content_type:

        return "json"


    if path.endswith(
        ".xlsx"
    ):

        return "xlsx"


    if path.endswith(
        ".xls"
    ):

        return "xls"


    if "spreadsheetml" in content_type:

        return "xlsx"


    if "excel" in content_type:

        return "xls"


    if "html" in content_type:

        return "html"


    return "unknown"


def get_extension_score(url):

    path = urlparse(
        url
    ).path.lower()

    if path.endswith(
        ".csv"
    ):
        return 60

    if path.endswith(
        ".json"
    ):
        return 55

    if path.endswith(
        ".xlsx"
    ):
        return 50

    if path.endswith(
        ".xls"
    ):
        return 35

    return 0


def read_response_bytes(
    response,
    max_bytes=MAX_DOWNLOAD_BYTES
):

    content_length = response.headers.get(
        "Content-Length"
    )


    if content_length:

        try:

            if int(
                content_length
            ) > max_bytes:

                return (
                    None,
                    True
                )

        except ValueError:

            pass


    chunks = []
    total_bytes = 0


    while True:

        chunk = response.read(
            1024 * 1024
        )

        if not chunk:

            break


        total_bytes += len(
            chunk
        )

        if total_bytes > max_bytes:

            return (
                None,
                True
            )


        chunks.append(
            chunk
        )


    return (
        b"".join(
            chunks
        ),
        False
    )


def fetch_url(
    url,
    timeout=10,
    max_bytes=MAX_DOWNLOAD_BYTES,
    urlopen_func=None
):

    if urlopen_func is None:

        urlopen_func = urlopen


    request = Request(
        url,
        headers={
            "User-Agent":
                "DataObservatory/1.0"
        }
    )


    try:

        with urlopen_func(
            request,
            timeout=timeout
        ) as response:

            data, too_large = read_response_bytes(
                response,
                max_bytes=max_bytes
            )

            return {
                "status":
                    "too_large" if too_large else "ok",

                "content":
                    data,

                "content_type":
                    response.headers.get(
                        "Content-Type",
                        ""
                    ).lower(),

                "final_url":
                    response.geturl()
            }


    except (HTTPError, URLError, TimeoutError, OSError) as error:

        return {
            "status":
                "failed",

            "message":
                str(
                    error
                )
        }


def parse_csv_bytes(content):

    text = content.decode(
        "utf-8-sig"
    )

    sample = text[:4096]

    try:

        dialect = csv.Sniffer().sniff(
            sample,
            delimiters=[
                ",",
                ";",
                "\t",
                "|"
            ]
        )

    except csv.Error:

        dialect = csv.excel


    reader = csv.DictReader(
        io.StringIO(
            text
        ),
        dialect=dialect
    )


    return [
        normalise_row(
            row
        )
        for row
        in reader
    ]


def parse_json_bytes(content):

    parsed = json.loads(
        content.decode(
            "utf-8-sig"
        )
    )


    if isinstance(
        parsed,
        list
    ):

        records = parsed


    elif isinstance(
        parsed,
        dict
    ):

        dbnomics_rows = parse_dbnomics_series_rows(
            parsed
        )

        if dbnomics_rows is not None:
            return dbnomics_rows

        json_stat_rows = parse_json_stat_rows(
            parsed
        )

        if json_stat_rows is not None:
            return json_stat_rows

        records = None

        for key in [
            "data",
            "results",
            "items"
        ]:

            if isinstance(
                parsed.get(
                    key
                ),
                list
            ):

                records = parsed[key]
                break


        if records is None:

            return None


    else:

        return None


    if not all(
        isinstance(
            item,
            dict
        )
        for item
        in records
    ):

        return None


    return [
        normalise_row(
            row
        )
        for row
        in records
    ]


def parse_json_stat_rows(parsed):

    if not isinstance(
        parsed,
        dict
    ):
        return None

    dimensions = parsed.get(
        "dimension"
    )
    ids = parsed.get(
        "id"
    )
    sizes = parsed.get(
        "size"
    )
    values = parsed.get(
        "value"
    )

    if not (
        isinstance(
            dimensions,
            dict
        )
        and
        isinstance(
            ids,
            list
        )
        and
        isinstance(
            sizes,
            list
        )
        and
        values is not None
    ):
        return None

    dimension_values = []

    for dimension_id in ids:
        dimension = dimensions.get(
            dimension_id,
            {}
        )
        category = dimension.get(
            "category",
            {}
        )
        index = category.get(
            "index",
            {}
        )
        labels = category.get(
            "label",
            {}
        )

        if isinstance(
            index,
            dict
        ):
            ordered_codes = [
                code
                for code, _
                in sorted(
                    index.items(),
                    key=lambda item: item[1]
                )
            ]
        elif isinstance(
            index,
            list
        ):
            ordered_codes = index
        else:
            return None

        dimension_values.append(
            [
                (
                    code,
                    labels.get(
                        code,
                        code
                    )
                    if isinstance(
                        labels,
                        dict
                    )
                    else
                    code
                )
                for code
                in ordered_codes
            ]
        )

    if not dimension_values:
        return None

    total_size = 1

    for size in sizes:
        total_size *= int(
            size
        )

    rows = []

    for flat_index in range(
        total_size
    ):
        remainder = flat_index
        row = {}

        for dimension_position in reversed(
            range(
                len(
                    ids
                )
            )
        ):
            size = int(
                sizes[dimension_position]
            )
            value_index = remainder % size
            remainder = remainder // size
            dimension_id = ids[dimension_position]
            code, label = dimension_values[dimension_position][value_index]
            row[dimension_id] = code
            row[dimension_id + "_label"] = label

        if isinstance(
            values,
            dict
        ):
            observation_value = values.get(
                str(
                    flat_index
                )
            )
        elif isinstance(
            values,
            list
        ):
            observation_value = (
                values[flat_index]
                if flat_index < len(
                    values
                )
                else
                None
            )
        else:
            observation_value = None

        if observation_value is None:
            continue

        row["value"] = observation_value
        rows.append(
            normalise_row(
                row
            )
        )

    return rows


def parse_dbnomics_series_rows(parsed):

    series_docs = (
        parsed.get(
            "series",
            {}
        ).get(
            "docs",
            []
        )
        if isinstance(
            parsed,
            dict
        )
        else
        []
    )

    if not series_docs:
        return None

    rows = []

    for series in series_docs:

        if not isinstance(
            series,
            dict
        ):
            continue

        periods = series.get(
            "period",
            []
        )
        values = series.get(
            "value",
            []
        )
        dimensions = series.get(
            "dimensions",
            {}
        )

        if not isinstance(
            dimensions,
            dict
        ):
            dimensions = {}

        for index, period in enumerate(
            periods
        ):

            value = (
                values[index]
                if index < len(
                    values
                )
                else
                None
            )

            row = {
                "dataset_code":
                    series.get(
                        "dataset_code"
                    ),

                "dataset_name":
                    series.get(
                        "dataset_name"
                    ),

                "series_code":
                    series.get(
                        "series_code"
                    ),

                "series_name":
                    series.get(
                        "series_name"
                    ),

                "time":
                    period,

                "value":
                    value
            }

            row.update(
                dimensions
            )
            rows.append(
                normalise_row(
                    row
                )
            )

    return rows


def parse_xlsx_bytes(content):

    try:

        from openpyxl import load_workbook

    except ImportError:

        return None


    workbook = load_workbook(
        io.BytesIO(
            content
        ),
        read_only=True,
        data_only=True
    )

    sheet = workbook.active

    rows = sheet.iter_rows(
        values_only=True
    )


    try:

        headers = next(
            rows
        )

    except StopIteration:

        return []


    headers = [
        str(
            header
        )
        if header is not None
        else f"column_{index + 1}"

        for index, header
        in enumerate(
            headers
        )
    ]


    parsed_rows = []

    for values in rows:

        parsed_rows.append(
            normalise_row(
                dict(
                    zip(
                        headers,
                        values
                    )
                )
            )
        )


    return parsed_rows


class TableParser(HTMLParser):

    def __init__(self):
        super().__init__()
        self.tables = []
        self.current_table = None
        self.current_row = None
        self.current_cell = None


    def handle_starttag(
        self,
        tag,
        attrs
    ):
        tag = tag.lower()

        if tag == "table":
            self.current_table = []

        elif tag == "tr" and self.current_table is not None:
            self.current_row = []

        elif tag in {
            "td",
            "th"
        } and self.current_row is not None:
            self.current_cell = []


    def handle_data(
        self,
        data
    ):
        if self.current_cell is not None:
            self.current_cell.append(
                data
            )


    def handle_endtag(
        self,
        tag
    ):
        tag = tag.lower()

        if tag in {
            "td",
            "th"
        } and self.current_cell is not None:
            self.current_row.append(
                " ".join(
                    " ".join(
                        self.current_cell
                    ).split()
                )
            )
            self.current_cell = None

        elif tag == "tr" and self.current_row is not None:
            if any(
                self.current_row
            ):
                self.current_table.append(
                    self.current_row
                )
            self.current_row = None

        elif tag == "table" and self.current_table is not None:
            if self.current_table:
                self.tables.append(
                    self.current_table
                )
            self.current_table = None


def parse_html_table_bytes(content):

    parser = TableParser()
    parser.feed(
        content.decode(
            "utf-8",
            errors="ignore"
        )
    )

    for table in parser.tables:
        if len(
            table
        ) < 2:
            continue

        headers = [
            header
            or
            f"column_{index + 1}"
            for index, header
            in enumerate(
                table[0]
            )
        ]

        rows = []

        for values in table[1:]:
            padded_values = values + [
                None
            ] * max(
                0,
                len(
                    headers
                )
                - len(
                    values
                )
            )

            rows.append(
                normalise_row(
                    dict(
                        zip(
                            headers,
                            padded_values
                        )
                    )
                )
            )

        if rows:
            return rows

    return None


def parse_rows(
    content,
    data_format
):

    if data_format == "csv":

        return parse_csv_bytes(
            content
        )


    if data_format == "json":

        return parse_json_bytes(
            content
        )


    if data_format == "xlsx":

        return parse_xlsx_bytes(
            content
        )

    if data_format == "html":

        return parse_html_table_bytes(
            content
        )


    return None


def find_data_links(
    html_content,
    base_url
):

    class LinkParser(HTMLParser):

        def __init__(self):

            super().__init__()
            self.links = []
            self.current_href = None
            self.current_text = []


        def handle_starttag(
            self,
            tag,
            attrs
        ):

            if tag.lower() != "a":

                return


            attrs = dict(
                attrs
            )

            href = attrs.get(
                "href"
            )

            if href:
                self.current_href = href
                self.current_text = []


        def handle_data(
            self,
            data
        ):
            if self.current_href:
                self.current_text.append(
                    data
                )


        def handle_endtag(
            self,
            tag
        ):
            if tag.lower() == "a" and self.current_href:
                self.links.append(
                    {
                        "href":
                            self.current_href,

                        "text":
                            " ".join(
                                " ".join(
                                    self.current_text
                                ).split()
                            )
                    }
                )
                self.current_href = None
                self.current_text = []


    parser = LinkParser()

    parser.feed(
        html_content.decode(
            "utf-8",
            errors="ignore"
        )
    )


    candidate_links = []

    for link in parser.links:

        absolute_url = urljoin(
            base_url,
            link.get(
                "href",
                ""
            )
        )

        path = urlparse(
            absolute_url
        ).path.lower()

        if any(
            path.endswith(
                extension
            )
            for extension
            in SUPPORTED_EXTENSIONS
        ):

            candidate_links.append(
                absolute_url
            )


    return candidate_links


def discover_data_resources(
    html_content,
    base_url,
    data_target=None,
    publisher=None
):

    class DiscoveryLinkParser(HTMLParser):

        def __init__(self):
            super().__init__()
            self.links = []
            self.current_href = None
            self.current_text = []


        def handle_starttag(
            self,
            tag,
            attrs
        ):
            if tag.lower() != "a":
                return

            attrs = dict(
                attrs
            )
            href = attrs.get(
                "href"
            )

            if href:
                self.current_href = href
                self.current_text = []


        def handle_data(
            self,
            data
        ):
            if self.current_href:
                self.current_text.append(
                    data
                )


        def handle_endtag(
            self,
            tag
        ):
            if tag.lower() == "a" and self.current_href:
                self.links.append(
                    {
                        "href":
                            self.current_href,

                        "text":
                            " ".join(
                                " ".join(
                                    self.current_text
                                ).split()
                            )
                    }
                )
                self.current_href = None
                self.current_text = []


    parser = DiscoveryLinkParser()
    parser.feed(
        html_content.decode(
            "utf-8",
            errors="ignore"
        )
    )

    base_host = urlparse(
        base_url
    ).hostname
    target_terms = set(
        re.findall(
            r"[a-z0-9]+",
            (
                data_target
                or
                ""
            ).lower()
        )
    )
    candidates = []

    for link in parser.links:
        absolute_url = urljoin(
            base_url,
            link.get(
                "href",
                ""
            )
        )
        text = link.get(
            "text",
            ""
        )
        haystack = (
            absolute_url
            + " "
            + text
        ).lower()
        path = urlparse(
            absolute_url
        ).path.lower()

        has_supported_extension = any(
            path.endswith(
                extension
            )
            for extension
            in SUPPORTED_EXTENSIONS
        )
        has_keyword = any(
            keyword in haystack
            for keyword
            in RESOURCE_KEYWORDS
        )

        if not (
            has_supported_extension
            or
            has_keyword
        ):
            continue

        score = get_extension_score(
            absolute_url
        )

        if base_host and urlparse(
            absolute_url
        ).hostname == base_host:
            score += 15

        if has_keyword:
            score += 10

        if publisher and publisher.lower() in haystack:
            score += 5

        score += 3 * len(
            target_terms.intersection(
                set(
                    re.findall(
                        r"[a-z0-9]+",
                        haystack
                    )
                )
            )
        )

        candidates.append(
            {
                "url":
                    absolute_url,

                "text":
                    text,

                "score":
                    score
            }
        )

    candidates.sort(
        key=lambda candidate: candidate["score"],
        reverse=True
    )

    return candidates


def build_result(
    retrieval_status,
    message=None,
    data_access_url=None,
    data_format=None,
    rows=None,
    retrieval_method=None,
    retrieved_scope=None,
    validation_status=None,
    warnings=None,
    error=None,
    metadata=None
):

    rows = rows or []

    stored_rows = rows[:MAX_STORED_ROWS]


    result = {
        "retrieval_status":
            retrieval_status,

        "message":
            message or "",

        "data_access_url":
            data_access_url,

        "format":
            data_format,

        "retrieval_method":
            retrieval_method,

        "retrieved_scope":
            retrieved_scope,

        "validation_status":
            validation_status,

        "warnings":
            warnings or [],

        "error":
            error,

        "metadata":
            metadata or {},

        "retrieved_row_count":
            len(
                rows
            )
            if rows is not None
            else None,

        "stored_row_count":
            len(
                stored_rows
            ),

        "rows":
            stored_rows
    }


    if rows and len(
        rows
    ) > MAX_STORED_ROWS:

        result["message"] = (
            f"Retrieved {len(rows)} rows. "
            f"Stored the first {MAX_STORED_ROWS} rows."
        )


    return result


def build_eurostat_api_url(dataset_code):

    normalised_code = normalize_eurostat_dataset_code(
        dataset_code
    )

    if not normalised_code:
        return None

    return EUROSTAT_JSONSTAT_ENDPOINT.format(
        code=normalised_code.upper()
    )


def validate_retrieved_rows(
    rows,
    data_target=None,
    metadata=None
):

    if not rows:
        return (
            "invalid",
            "Retrieved resource did not contain parseable observations."
        )

    if (
        metadata
        and
        metadata.get(
            "publisher"
        ) == "Eurostat"
        and
        metadata.get(
            "dataset_code"
        )
    ):
        return (
            "validated",
            ""
        )

    if not data_target:
        return (
            "validated",
            ""
        )

    target_terms = {
        term
        for term
        in re.findall(
            r"[a-z0-9]+",
            data_target.lower()
        )
        if len(
            term
        ) > 2
    }

    if not target_terms:
        return (
            "validated",
            ""
        )

    sample_text = " ".join(
        [
            " ".join(
                str(
                    key
                )
                + " "
                + str(
                    value
                )
                for key, value
                in row.items()
            )
            for row
            in rows[:20]
        ]
        +
        [
            json.dumps(
                metadata or {},
                ensure_ascii=False
            )
        ]
    ).lower()

    if any(
        term in sample_text
        for term
        in target_terms
    ):
        return (
            "validated",
            ""
        )

    return (
        "needs_review",
        "Retrieved rows were parsed, but relevance to the Data Target could not be confirmed from headers or metadata."
    )


def retrieve_parsed_resource(
    resource_url,
    data_target=None,
    publisher=None,
    retrieval_method=None,
    retrieved_scope=None,
    metadata=None,
    urlopen_func=None,
    max_bytes=MAX_DOWNLOAD_BYTES
):

    fetched = fetch_url(
        resource_url,
        max_bytes=max_bytes,
        urlopen_func=urlopen_func
    )

    if fetched["status"] == "too_large":
        return build_result(
            "too_large",
            message="The source is larger than the allowed download size.",
            data_access_url=resource_url,
            retrieval_method=retrieval_method,
            retrieved_scope=retrieved_scope,
            validation_status="invalid",
            metadata=metadata
        )

    if fetched["status"] == "failed":
        return build_result(
            "failed",
            message=fetched.get(
                "message",
                "The source could not be retrieved."
            ),
            data_access_url=resource_url,
            retrieval_method=retrieval_method,
            retrieved_scope=retrieved_scope,
            validation_status="invalid",
            error=fetched.get(
                "message"
            ),
            metadata=metadata
        )

    final_url = fetched.get(
        "final_url",
        resource_url
    )
    data_format = detect_format(
        final_url,
        fetched.get(
            "content_type",
            ""
        )
    )
    rows = parse_rows(
        fetched["content"],
        data_format
    )

    if rows is None:
        return build_result(
            "unsupported",
            message="The source format could not be interpreted as tabular records.",
            data_access_url=final_url,
            data_format=data_format,
            retrieval_method=retrieval_method,
            retrieved_scope=retrieved_scope,
            validation_status="invalid",
            metadata=metadata
        )

    validation_status, warning = validate_retrieved_rows(
        rows,
        data_target=data_target,
        metadata={
            "publisher":
                publisher,

            "resource_url":
                final_url,

            **(
                metadata
                or
                {}
            )
        }
    )

    if validation_status == "invalid":
        return build_result(
            "unsupported",
            message=warning,
            data_access_url=final_url,
            data_format=data_format,
            rows=rows,
            retrieval_method=retrieval_method,
            retrieved_scope=retrieved_scope,
            validation_status=validation_status,
            warnings=[
                warning
            ]
            if warning
            else
            [],
            metadata={
                "publisher":
                    publisher,

                **(
                    metadata
                    or
                    {}
                )
            }
        )

    return build_result(
        "retrieved"
        if validation_status == "validated"
        else
        "retrieved_with_warnings",
        message=warning,
        data_access_url=final_url,
        data_format=data_format,
        rows=rows,
        retrieval_method=retrieval_method,
        retrieved_scope=retrieved_scope,
        validation_status=validation_status,
        warnings=[
            warning
        ]
        if warning
        else
        [],
        metadata={
            "publisher":
                publisher,

            **(
                metadata
                or
                {}
            )
        }
    )


def retrieve_dataset(
    source_url,
    data_target=None,
    publisher=None,
    dataset_code=None,
    original_source_url=None,
    urlopen_func=None,
    max_bytes=MAX_DOWNLOAD_BYTES
):

    original_source_url = original_source_url or source_url
    eurostat_code = (
        dataset_code
        or
        extract_eurostat_dataset_code(
            source_url
        )
        or
        extract_eurostat_dataset_code(
            original_source_url
        )
    )

    if eurostat_code:
        eurostat_resource_url = build_eurostat_api_url(
            eurostat_code
        )
        custom_view = detect_eurostat_custom_view(
            original_source_url
        )
        eurostat_result = retrieve_parsed_resource(
            eurostat_resource_url,
            data_target=data_target,
            publisher=publisher or "Eurostat",
            retrieval_method="official_api",
            retrieved_scope="dataset",
            metadata={
                "dataset_code":
                    normalize_eurostat_dataset_code(
                        eurostat_code
                    ),

                "source_url":
                    original_source_url,

                "custom_selection_resolved":
                    False,

                "is_custom_view":
                    custom_view.get(
                        "is_custom_view",
                        False
                    )
            },
            urlopen_func=urlopen_func,
            max_bytes=max_bytes
        )

        if eurostat_result.get(
            "retrieval_status"
        ) in {
            "retrieved",
            "retrieved_with_warnings"
        }:
            return eurostat_result

    first_fetch = fetch_url(
        source_url,
        max_bytes=max_bytes,
        urlopen_func=urlopen_func
    )


    if first_fetch["status"] == "too_large":

        return build_result(
            "too_large",
            message="The source is larger than the allowed download size.",
            data_access_url=source_url,
            validation_status="invalid"
        )


    if first_fetch["status"] == "failed":

        return build_result(
            "failed",
            message=first_fetch.get(
                "message",
                "The source could not be retrieved."
            ),
            data_access_url=source_url,
            validation_status="invalid",
            error=first_fetch.get(
                "message"
            )
        )


    final_url = first_fetch.get(
        "final_url",
        source_url
    )

    data_format = detect_format(
        final_url,
        first_fetch.get(
            "content_type",
            ""
        )
    )


    if data_format == "html":

        rows = parse_rows(
            first_fetch["content"],
            data_format
        )

        if rows is not None:
            validation_status, warning = validate_retrieved_rows(
                rows,
                data_target=data_target,
                metadata={
                    "publisher":
                        publisher,

                    "source_url":
                        original_source_url
                }
            )

            if validation_status == "invalid":
                return build_result(
                    "unsupported",
                    message=warning,
                    data_access_url=final_url,
                    data_format=data_format,
                    rows=rows,
                    retrieval_method="html_table",
                    retrieved_scope="dataset",
                    validation_status=validation_status,
                    warnings=[
                        warning
                    ]
                    if warning
                    else
                    [],
                    metadata={
                        "source_url":
                            original_source_url
                    }
                )

            return build_result(
                "retrieved"
                if validation_status == "validated"
                else
                "retrieved_with_warnings",
                message=warning,
                data_access_url=final_url,
                data_format=data_format,
                rows=rows,
                retrieval_method="html_table",
                retrieved_scope="dataset",
                validation_status=validation_status,
                warnings=[
                    warning
                ]
                if warning
                else
                [],
                metadata={
                    "source_url":
                        original_source_url
                }
            )

        candidates = discover_data_resources(
            first_fetch["content"],
            final_url,
            data_target=data_target,
            publisher=publisher
        )

        for candidate in candidates[:10]:
            candidate_result = retrieve_parsed_resource(
                candidate["url"],
                data_target=data_target,
                publisher=publisher,
                retrieval_method="resource_discovery",
                retrieved_scope="dataset",
                metadata={
                    "source_url":
                        original_source_url,

                    "resource_link_text":
                        candidate.get(
                            "text"
                        ),

                    "resource_score":
                        candidate.get(
                            "score"
                        )
                },
                urlopen_func=urlopen_func,
                max_bytes=max_bytes
            )

            if candidate_result.get(
                "retrieval_status"
            ) in {
                "retrieved",
                "retrieved_with_warnings"
            }:
                return candidate_result


        return build_result(
            "unsupported",
            message="No reliable structured data resource was found after inspecting the landing page.",
            data_access_url=source_url,
            retrieval_method="resource_discovery",
            retrieved_scope="dataset",
            validation_status="invalid",
            metadata={
                "source_url":
                    original_source_url,

                "candidate_count":
                    len(
                        candidates
                    )
            }
        )


    if data_format == "xls":

        return build_result(
            "unsupported",
            message="Legacy .xls files are not supported in this MVP.",
            data_access_url=final_url,
            data_format=data_format,
            retrieval_method="direct_download",
            retrieved_scope="dataset",
            validation_status="invalid"
        )


    rows = parse_rows(
        first_fetch["content"],
        data_format
    )

    if rows is None:

        return build_result(
            "unsupported",
            message="The source format could not be interpreted as tabular records.",
            data_access_url=final_url,
            data_format=data_format,
            retrieval_method="direct_download",
            retrieved_scope="dataset",
            validation_status="invalid"
        )

    validation_status, warning = validate_retrieved_rows(
        rows,
        data_target=data_target,
        metadata={
            "publisher":
                publisher,

            "source_url":
                original_source_url
        }
    )

    if validation_status == "invalid":
        return build_result(
            "unsupported",
            message=warning,
            data_access_url=final_url,
            data_format=data_format,
            rows=rows,
            retrieval_method="direct_download",
            retrieved_scope="dataset",
            validation_status=validation_status,
            warnings=[
                warning
            ]
            if warning
            else
            [],
            metadata={
                "source_url":
                    original_source_url
            }
        )

    return build_result(
        "retrieved"
        if validation_status == "validated"
        else
        "retrieved_with_warnings",
        message=warning,
        data_access_url=final_url,
        data_format=data_format,
        rows=rows,
        retrieval_method="direct_download",
        retrieved_scope="dataset",
        validation_status=validation_status,
        warnings=[
            warning
        ]
        if warning
        else
        [],
        metadata={
            "source_url":
                original_source_url
        }
    )
